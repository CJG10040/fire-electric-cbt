"""Write the review workbook (duplicate clusters, per-item mapping, topic counts)."""
import pickle
import re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sa_final import TOPIC_ORDER

items, clusters, meta = pickle.load(open("sa_final.pkl", "rb"))
M = {i["id"]: i for i in items}


def label(it):
    blk = f"{it['y']}년" + (f" {it['r']}회" if it["r"] else "")
    return f"{blk} #{it['n']}"


def plain(t, n=160):
    return re.sub(r"\s+", " ", t)[:n]


wb = Workbook()
ws = wb.active
ws.title = "중복 묶음"
ws.append(["묶음ID", "주제", "대표 문항(최신)", "출제 횟수", "출제 이력(최신→과거)", "답 표기 차이", "차이/주의 메모", "대표 정답(요약)"])
for c in clusters:
    rep = M[c["rep"]]
    hist = " · ".join(label(M[i]) for i in c["m"])
    memo = c.get("nt") or ("출제본마다 숫자·항목이 다름 — 앱의 '출제 이력'에서 비교" if c["ad"] else "")
    ws.append([c["k"], TOPIC_ORDER[c["tp"]], rep["t"], len(c["m"]), hist, "있음" if c["ad"] else "", memo, plain(rep["at"] if "at" in rep else re.sub("<[^>]+>", " ", rep["ah"]))])

ws2 = wb.create_sheet("문항별(447)")
ws2.append(["문항ID", "출제", "번호", "제목", "유형", "주제", "묶음ID", "묶음 내 출제 횟수", "대표 여부", "대표와 질문 차이", "대표와 답 차이"])
for it in items:
    c = next(x for x in clusters if x["k"] == it["c"])
    ws2.append([it["id"], label(it).rsplit(" #", 1)[0], it["n"], it["t"], it["kind"] + (f"({it['pick']}가지)" if it["pick"] else ""), it["topic"], it["c"], len(c["m"]), "대표" if c["rep"] == it["id"] else "", "차이" if it["dq"] else "", "차이" if it["da"] else ""])

ws3 = wb.create_sheet("주제별 문항 수")
ws3.append(["주제", "고유 문항(묶음)", "원본 출제건수"])
for ti, t in enumerate(TOPIC_ORDER):
    cl = [c for c in clusters if c["tp"] == ti]
    ws3.append([t, len(cl), sum(len(c["m"]) for c in cl)])
ws3.append(["합계", len(clusters), sum(len(c["m"]) for c in clusters)])

ws4 = wb.create_sheet("읽는 법")
for line in [
    "출제 횟수 = 같은 질문 의도·답 핵심이 반복된 원본 문항 건수(연도 수가 아님). 같은 시험 블록 안의 두 문항은 합치지 않았습니다.",
    "대표 문항 = 묶음 중 가장 최신 출제본. 학습 화면(전체·주제별)에는 대표 1문항만 출제하고, 연도별 모드는 그 해의 원본 문항을 그대로 출제합니다.",
    "자동 유사도 묶기 후 사람이 검수해 21개 묶음을 병합하고 1건을 분리했습니다. 질문 형태가 다른 것(예: '장점' vs '장·단점', 설치대상 vs 설치기준)은 합치지 않았습니다.",
    "출처는 암기영상 정리본(비공식)이며 정답은 원문 그대로입니다. 연도에 따라 기준이 개정되었을 수 있으니 확인이 필요한 항목은 '답 표기 차이'로 표시했습니다.",
    "2010~2023년은 정리본에 회차 구분이 없어 연도 단위로 표시합니다. 2025년 시험은 정리본에 없으므로 출제 횟수에 포함되지 않았습니다.",
]:
    ws4.append([line])

head_fill = PatternFill("solid", fgColor="D9E2F3")
for w in (ws, ws2, ws3):
    for cell in w[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
    w.freeze_panes = "A2"
for w, widths in ((ws, [14, 22, 46, 9, 60, 10, 50, 50]), (ws2, [16, 12, 6, 46, 10, 22, 14, 10, 8, 10, 10]), (ws3, [26, 16, 14]), (ws4, [150])):
    for i, wd in enumerate(widths, 1):
        w.column_dimensions[get_column_letter(i)].width = wd
    for row in w.iter_rows(min_row=1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
wb.save("단답서술_중복묶음_검토표.xlsx")
print("saved", len(clusters), "clusters;", len(items), "items")
